"""Build the 3-slide deck docs/slides.pptx (16:9) following docs/DEMO_SCRIPT.md §A.

Every number is read from results/*.json, so the deck cannot drift from the results.
Usage: python scripts/make_slides.py
"""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
OUT = ROOT / "docs" / "slides.pptx"
HW_JOB = "db3b1bimb58s7387e0jg"
HW_HIST = ROOT / "results" / "hardware" / "ibm_job_histogram.png"
HIST_CAPTION = (
    "Real ibm_kingston output — tallest bar 1001 = optimal basket (RELIANCE+ITC), "
    "second 1100 = runner-up"
)

PURPLE = RGBColor(0x3B, 0x1F, 0x7A)
GREY = RGBColor(0x44, 0x44, 0x44)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xE8, 0x67, 0x38)
REPO = "github.com/achintya007-arch/qfolio"
REPORT = "achintya007-arch.github.io/qfolio"


def load_numbers() -> dict:
    """Pull the handful of numbers the slides quote from the results JSON."""
    bench = json.loads((ROOT / "results" / "benchmark.json").read_text())
    hw = json.loads((ROOT / "results" / "hardware" / f"{HW_JOB}.json").read_text())
    s = bench["summary"]
    run2 = next(r for r in hw["runs"] if r["reps"] == 2)
    return {
        "xy": s["xy_qaoa_p2"]["p_optimal_mean"],
        "rand": s["random"]["p_optimal_mean"],
        "pen": s["penalty_qaoa_p2"]["p_optimal_mean"],
        "pen_valid": s["penalty_qaoa_p2"]["p_feasible_mean"],
        "n_inst": s["xy_qaoa_p2"]["instances"],
        "backend": hw["backend"],
        "hw_raw": run2["raw"]["p_optimal"],
        "hw_ps": run2["postselected"]["p_optimal"],
        "hw_rand": hw["p_random"],
    }


def add_text(slide, x, y, w, h, lines, size=20, color=GREY, bold=False):
    """Add a text box; ``lines`` is a list of strings (one paragraph each)."""
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.space_after = Pt(6)
    return tf


def add_title(slide, title, subtitle=None):
    add_text(slide, 0.5, 0.3, 12.3, 0.9, [title], size=34, color=PURPLE, bold=True)
    if subtitle:
        add_text(slide, 0.5, 1.1, 12.3, 0.5, [subtitle], size=18)


def add_box(slide, x, y, w, h, title, body, color):
    """A simple coloured 'circuit cartoon' card: heading plus two lines."""
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))  # 1 = rectangle
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xF7, 0xF6, 0xFA)
    box.line.color.rgb = color
    box.line.width = Pt(2.5)
    add_text(slide, x + 0.2, y + 0.1, w - 0.4, 0.5, [title], size=22, color=color, bold=True)
    add_text(slide, x + 0.2, y + 0.65, w - 0.4, h - 0.7, body, size=17)


def build(n: dict) -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1: the problem.
    s1 = prs.slides.add_slide(blank)
    add_title(
        s1,
        "Picking the best 3 of 6 stocks, on a quantum computer",
        "Robo-advisor model baskets · NSE large caps · return ↑ risk ↓",
    )
    s1.shapes.add_picture(str(FIG / "frontier.png"), Inches(0.5), Inches(1.75), height=Inches(5.5))
    add_text(
        s1,
        9.2,
        2.2,
        3.8,
        5.0,
        [
            "All C(6,3) = 20 baskets on the risk/return plane.",
            "Basket count grows as C(n,k): ~185 000 for 10-of-20.",
            "Goal: find the best basket for a client's risk aversion q.",
        ],
        size=20,
    )

    # Slide 2: the idea and the result.
    s2 = prs.slides.add_slide(blank)
    add_title(s2, "Put the constraint in the circuit, not in the cost")
    add_box(
        s2,
        0.5,
        1.25,
        6.0,
        1.6,
        "Penalty QAOA (X mixer)",
        [
            "Searches all 64 bit strings, 44 of them invalid.",
            "→ finds the best basket at chance level",
        ],
        ORANGE,
    )
    add_box(
        s2,
        6.85,
        1.25,
        6.0,
        1.6,
        "XY-QAOA (Dicke state + XY mixer)",
        ["Only ever visits the 20 valid baskets.", "→ every shot is a real portfolio"],
        BLUE,
    )
    s2.shapes.add_picture(str(FIG / "headline.png"), Inches(0.5), Inches(3.05), width=Inches(12.3))
    ratio = n["xy"] / n["rand"]
    add_text(
        s2,
        0.5,
        7.0,
        12.3,
        0.45,
        [
            f"~{ratio:.0f}× chance (P(optimal) {n['xy']:.2f} vs {n['rand']:.2f}, "
            f"{n['n_inst']} instances) · 100% valid · penalty-QAOA = chance "
            f"({n['pen']:.2f}, {n['pen_valid']:.0%} valid)"
        ],
        size=18,
        color=PURPLE,
        bold=True,
    )

    # Slide 3: real hardware and honesty.
    s3 = prs.slides.add_slide(blank)
    add_title(
        s3,
        f"Real hardware ({n['backend']}) and honest limits",
        f"P(optimal) {n['hw_raw']:.2f} raw → {n['hw_ps']:.2f} post-selected "
        f"vs {n['hw_rand']:.2f} random · job {HW_JOB}",
    )
    # Left: simulated + real noise ladder. Right: the raw IBM dashboard histogram as evidence.
    s3.shapes.add_picture(
        str(FIG / "noise_ladder.png"), Inches(0.4), Inches(1.7), width=Inches(6.4)
    )
    s3.shapes.add_picture(str(HW_HIST), Inches(6.95), Inches(1.75), width=Inches(6.0))
    add_text(s3, 6.95, 4.4, 6.0, 0.8, [HIST_CAPTION], size=14, color=PURPLE)
    add_text(
        s3,
        0.5,
        5.2,
        8.6,
        2.0,
        [
            "• Symmetry post-selection = free error detection",
            "• Classical brute force is still faster at this size: no quantum advantage claimed",
            "• Next: more business constraints as circuit symmetries",
        ],
        size=20,
    )
    add_text(s3, 9.3, 6.2, 3.8, 1.0, [f"Repo: {REPO}", f"Report: {REPORT}"], size=14, color=PURPLE)
    return prs


def main() -> None:
    build(load_numbers()).save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
